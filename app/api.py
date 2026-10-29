from datetime import date
from difflib import SequenceMatcher
from io import BytesIO
import json
import math
import mimetypes
import re
from urllib.parse import quote
from fastapi import APIRouter, Depends, HTTPException, UploadFile, File, Form
from fastapi.responses import Response
from fastapi.security import OAuth2PasswordRequestForm
from sqlalchemy import func
from sqlalchemy.orm import Session
from openpyxl import load_workbook
from .database import get_db
from .models import User, Employee, KPI, Entry, EmployeeKPI
from .schemas import Token, EmployeeSignup, EmployeeCreate, EmployeeOut, EntryUpdate, ReviewUpdate, GenerateRequest, EntryOut, KPIUpdate, KPICreate, KPIOut, EmployeeKPIAssignmentUpdate, UserRoleUpdate, UserOut
from .auth import verify_password, hash_password, create_token, current_user, require_roles
from .services import calculate, generate_month

router = APIRouter(prefix="/api")
CONTEXT_MAGIC = b"HLBCTX1\n"
MAX_CONTEXT_SIZE = 10 * 1024 * 1024
ALLOWED_CONTEXT_EXTENSIONS = {".pdf", ".doc", ".docx", ".xls", ".xlsx", ".ppt", ".pptx", ".txt", ".csv", ".png", ".jpg", ".jpeg"}
MAX_IMPORT_SIZE = 10 * 1024 * 1024

def _normalized_kpi_name(value) -> str:
    text = str(value or "").strip().lower()
    text = re.sub(r"^[a-z]\s*[.)]\s*", "", text)
    return re.sub(r"[^a-z0-9]+", " ", text).strip()

def _parse_excel_entry_rows(content: bytes) -> list[dict]:
    try:
        workbook = load_workbook(BytesIO(content), read_only=True, data_only=True)
    except Exception as exc:
        raise HTTPException(400, "Could not read the workbook. Upload a valid .xlsx file.") from exc
    try:
        for sheet in workbook.worksheets:
            for row_index, values in enumerate(sheet.iter_rows(values_only=True), start=1):
                headers = [_normalized_kpi_name(value) for value in values]
                kpi_col = next((i for i, value in enumerate(headers)
                                if value == "kpi" or value.startswith("kpi name") or value.startswith("description") or "key performance indicator" in value), None)
                id_col = next((i for i, value in enumerate(headers)
                               if value in ("kpi id", "kpi code", "kpi_id", "code")), None)
                actual_col = next((i for i, value in enumerate(headers) if "actual" in value), None)
                comments_col = next((i for i, value in enumerate(headers) if "comment" in value), None)
                if (kpi_col is not None or id_col is not None) and actual_col is not None:
                    parsed = []
                    for excel_row, data in enumerate(sheet.iter_rows(min_row=row_index + 1, values_only=True), start=row_index + 1):
                        label = data[kpi_col] if kpi_col is not None and kpi_col < len(data) else None
                        code = data[id_col] if id_col is not None and id_col < len(data) else None
                        actual = data[actual_col] if actual_col < len(data) else None
                        comments = data[comments_col] if comments_col is not None and comments_col < len(data) else None
                        if (label is not None or code is not None) and actual is not None and str(actual).strip() != "":
                            parsed.append({"label": str(label).strip() if label is not None else None,
                                           "code": str(code).strip() if code is not None else None,
                                           "actual": actual,
                                           "comments": str(comments).strip() if comments is not None else None,
                                           "row_number": excel_row})
                    if parsed:
                        return parsed
        raise HTTPException(422, "No KPI rows with Actual values were found. Include KPI name or KPI ID, Actual, and optional Comments columns.")
    finally:
        workbook.close()

def _match_import_kpi(row: dict, assignments: list[tuple[EmployeeKPI, KPI]]) -> KPI | None:
    if row.get("code"):
        raw_code = row["code"].strip()
        try:
            numeric_code = float(raw_code)
            if numeric_code.is_integer():
                raw_code = str(int(numeric_code))
        except ValueError:
            pass
        for _, kpi in assignments:
            if raw_code.casefold() == kpi.code.casefold() or raw_code == str(kpi.id):
                return kpi
    normalized = _normalized_kpi_name(row.get("label"))
    if not normalized:
        return None
    for _, kpi in assignments:
        if normalized == _normalized_kpi_name(kpi.name):
            return kpi
    candidate_scores = []
    tokens = set(normalized.split())
    for _, kpi in assignments:
        candidate = _normalized_kpi_name(kpi.name)
        candidate_tokens = set(candidate.split())
        overlap = len(tokens & candidate_tokens) / max(1, min(len(tokens), len(candidate_tokens)))
        score = max(SequenceMatcher(None, normalized, candidate).ratio(), overlap)
        candidate_scores.append((score, kpi))
    candidate_scores.sort(key=lambda item: item[0], reverse=True)
    if candidate_scores and candidate_scores[0][0] >= 0.58 and (len(candidate_scores) == 1 or candidate_scores[0][0] - candidate_scores[1][0] >= 0.08):
        return candidate_scores[0][1]
    return None

def _pack_context(filename: str, content_type: str, content: bytes) -> bytes:
    metadata = json.dumps({"filename": filename, "content_type": content_type}, ensure_ascii=False).encode("utf-8")
    return CONTEXT_MAGIC + len(metadata).to_bytes(4, "big") + metadata + content

def _unpack_context(value: bytes) -> tuple[str, str, bytes]:
    if value.startswith(CONTEXT_MAGIC) and len(value) >= len(CONTEXT_MAGIC) + 4:
        start = len(CONTEXT_MAGIC)
        header_size = int.from_bytes(value[start:start + 4], "big")
        header_start = start + 4
        header = json.loads(value[header_start:header_start + header_size].decode("utf-8"))
        return header["filename"], header["content_type"], value[header_start + header_size:]
    return "supporting-document", "application/octet-stream", value

def _check_entry_access(entry: Entry, user: User, db: Session):
    if user.role == "Employee":
        employee = db.query(Employee).filter(func.lower(Employee.email) == user.email.lower()).first()
        if not employee or employee.id != entry.employee_id:
            raise HTTPException(404, "Entry not found")
    elif user.role not in ("Admin", "Manager", "Partner"):
        raise HTTPException(403, "Insufficient permissions")

@router.post("/auth/signup", response_model=Token, status_code=201)
def signup(data: EmployeeSignup, db: Session = Depends(get_db)):
    email = data.email.lower()
    if db.query(User).filter(func.lower(User.email) == email).first():
        raise HTTPException(409, "An account with this email already exists")
    employee = db.query(Employee).filter(func.lower(Employee.email) == email, Employee.is_active.is_(True)).first()
    if not employee:
        raise HTTPException(403, "Your work email was not found. Ask your administrator to add your employee record first.")
    user = User(email=email, full_name=data.full_name.strip(), password_hash=hash_password(data.password), role="Employee", company=employee.team)
    db.add(user)
    db.commit()
    return Token(access_token=create_token(user.email))

@router.post("/auth/token", response_model=Token)
def login(form: OAuth2PasswordRequestForm = Depends(), db: Session = Depends(get_db)):
    user = db.query(User).filter(func.lower(User.email) == form.username.lower()).first()
    if not user or not verify_password(form.password, user.password_hash): raise HTTPException(401, "Incorrect credentials")
    return Token(access_token=create_token(user.email))

@router.get("/auth/me")
def me(user: User = Depends(current_user)):
    return {"id": user.id, "email": user.email, "full_name": user.full_name, "role": user.role}

@router.get("/profile")
def employee_profile(user: User = Depends(current_user), db: Session = Depends(get_db)):
    employee = db.query(Employee).filter(func.lower(Employee.email) == user.email.lower()).first()
    if not employee:
        raise HTTPException(404, "No employee profile is linked to this account")
    total_score, actual_sum = (db.query(func.sum(Entry.final_score), func.sum(Entry.actual))
                               .filter(Entry.employee_id == employee.id).one())
    rows = (db.query(Entry.month, func.sum(Entry.final_score), func.sum(Entry.actual))
            .filter(Entry.employee_id == employee.id)
            .group_by(Entry.month).order_by(Entry.month).all())
    return {"employee": {"id": employee.id, "name": employee.name, "email": employee.email,
                          "team": employee.team, "service_line": employee.service_line},
            "total_score": total_score or 0, "actual_sum": actual_sum or 0,
            "points": [{"month": month.isoformat(), "total_score": score or 0, "actual_sum": actual or 0}
                       for month, score, actual in rows]}

@router.get("/employees", response_model=list[EmployeeOut])
def employees(db: Session=Depends(get_db),
               _=Depends(current_user)
               ): return db.query(Employee).filter(Employee.is_active.is_(True)).all()

@router.get("/admin/users", response_model=list[UserOut])
def list_users(db: Session=Depends(get_db), _=Depends(require_roles("Admin"))):
#def list_users(db: Session=Depends(get_db), _=Depends(current_user)):
    return db.query(User).order_by(User.full_name, User.email).all()

@router.get("/reviewers", response_model=list[UserOut])
def list_reviewers(db: Session=Depends(get_db), _=Depends(require_roles("Admin", "Partner"))):
    return (db.query(User).filter(User.is_active.is_(True), User.role.in_(("Partner", "Manager")))
            .order_by(User.full_name, User.email).all())

@router.patch("/admin/users/{user_id}/role", response_model=UserOut)
def update_user_role(user_id: int, data: UserRoleUpdate, db: Session=Depends(get_db), admin: User=Depends(require_roles("Admin"))):
    user = db.get(User, user_id)
    if not user:
        raise HTTPException(404, "User not found")
    if user.id == admin.id and data.role != "Admin":
        raise HTTPException(400, "You cannot change your own role")
    if user.role == "Admin" and data.role != "Admin":
        active_admins = db.query(User).filter(User.role == "Admin", User.is_active.is_(True)).count()
        if active_admins <= 1:
            raise HTTPException(409, "At least one active Admin account must remain")
    user.role = data.role
    db.commit()
    db.refresh(user)
    return user

@router.delete("/admin/users/{user_id}", status_code=204)
def delete_user(user_id: int, db: Session=Depends(get_db), admin: User=Depends(require_roles("Admin"))):
    user = db.get(User, user_id)
    if not user:
        raise HTTPException(404, "User not found")
    if user.id == admin.id:
        raise HTTPException(400, "You cannot delete your own account")
    if user.role == "Admin" and user.is_active:
        active_admins = db.query(User).filter(User.role == "Admin", User.is_active.is_(True)).count()
        if active_admins <= 1:
            raise HTTPException(409, "At least one active Admin account must remain")
    db.delete(user)
    db.commit()
    return Response(status_code=204)

@router.get("/kpis", response_model=list[KPIOut])
# def list_kpis(db: Session=Depends(get_db), _=Depends(require_roles("Admin"))):
def list_kpis(db: Session=Depends(get_db)):
    return db.query(KPI).order_by(KPI.code).all()

@router.post("/kpis", response_model=KPIOut, status_code=201)
def create_kpi(data: KPICreate, db: Session=Depends(get_db), _=Depends(require_roles("Admin"))):
    code = data.code.strip()
    name = data.name.strip()
    perspective = data.perspective.strip()
    if not code or not name or not perspective:
        raise HTTPException(422, "KPI code, name, and perspective cannot be blank")
    if db.query(KPI).filter(func.lower(KPI.code) == code.lower()).first():
        raise HTTPException(409, "A KPI with this code already exists")
    kpi = KPI(code=code, name=name, perspective=perspective,
              perspective_weight=data.perspective_weight, kpi_weight=data.kpi_weight,
              target=data.target, unit_notes=data.unit_notes.strip() or None if data.unit_notes else None)
    db.add(kpi)
    db.commit()
    db.refresh(kpi)
    return kpi

def _employee_kpi_assignments(db: Session, employee_id: int):
    rows = (db.query(EmployeeKPI, KPI).join(KPI, EmployeeKPI.kpi_id == KPI.id)
            .filter(EmployeeKPI.employee_id == employee_id, EmployeeKPI.is_active.is_(True))
            .order_by(KPI.code).all())
    return [{"employee_id": assignment.employee_id, "kpi_id": kpi.id, "code": kpi.code,
             "name": kpi.name, "perspective": kpi.perspective,
             "perspective_weight": kpi.perspective_weight, "kpi_weight": kpi.kpi_weight,
             "target": assignment.target} for assignment, kpi in rows]

@router.get("/admin/employees/{employee_id}/kpis")
def get_employee_kpis(employee_id: int, db: Session=Depends(get_db), _=Depends(require_roles("Admin"))):
    employee = db.get(Employee, employee_id)
    if not employee or not employee.is_active:
        raise HTTPException(404, "Employee not found")
    return _employee_kpi_assignments(db, employee_id)

@router.put("/admin/employees/{employee_id}/kpis")
def set_employee_kpis(employee_id: int, data: EmployeeKPIAssignmentUpdate, db: Session=Depends(get_db), _=Depends(require_roles("Admin"))):
    employee = db.get(Employee, employee_id)
    if not employee or not employee.is_active:
        raise HTTPException(404, "Employee not found")
    ids = [assignment.kpi_id for assignment in data.assignments]
    if len(ids) != len(set(ids)):
        raise HTTPException(422, "Each KPI can only be assigned once")
    kpis = {kpi.id: kpi for kpi in db.query(KPI).filter(KPI.id.in_(ids)).all()} if ids else {}
    if len(kpis) != len(ids):
        raise HTTPException(422, "One or more KPI IDs do not exist")
    for assignment in db.query(EmployeeKPI).filter(EmployeeKPI.employee_id == employee_id).all():
        assignment.is_active = False
    for item in data.assignments:
        if item.target < 0:
            raise HTTPException(422, "Targets cannot be negative")
        assignment = db.get(EmployeeKPI, (employee_id, item.kpi_id))
        if assignment is None:
            assignment = EmployeeKPI(employee_id=employee_id, kpi_id=item.kpi_id, target=item.target, is_active=True)
            db.add(assignment)
        else:
            changed = assignment.target != item.target
            assignment.target = item.target
            assignment.is_active = True
            if changed:
                for entry in db.query(Entry).filter(
                    Entry.employee_id == employee_id, Entry.kpi_id == item.kpi_id,
                    Entry.status.in_(("Draft", "Rejected"))
                ).all():
                    entry.target_snapshot = item.target
                    calculate(entry, kpis[item.kpi_id])
    db.commit()
    return _employee_kpi_assignments(db, employee_id)

@router.patch("/kpis/{kpi_id}", response_model=KPIOut)
def update_kpi(kpi_id: int, data: KPIUpdate, db: Session=Depends(get_db), _=Depends(require_roles("Admin"))):
    kpi = db.get(KPI, kpi_id)
    if not kpi:
        raise HTTPException(404, "KPI not found")
    changes = data.model_dump(exclude_unset=True)
    if not changes:
        raise HTTPException(400, "No KPI changes provided")
    if any(value is None for field, value in changes.items() if field != "unit_notes"):
        raise HTTPException(422, "KPI name, perspective, weights, and target cannot be empty")
    for field in ("name", "perspective", "unit_notes"):
        if field in changes and changes[field] is not None:
            changes[field] = changes[field].strip() or (None if field == "unit_notes" else "")
    if not changes.get("name", kpi.name).strip() or not changes.get("perspective", kpi.perspective).strip():
        raise HTTPException(422, "KPI name and perspective cannot be blank")
    for field, value in changes.items():
        setattr(kpi, field, value)
    for entry in db.query(Entry).filter(Entry.kpi_id == kpi.id).all():
        calculate(entry, kpi)
    db.commit()
    db.refresh(kpi)
    return kpi

@router.get("/my/entries")
def my_entries(month: date|None=None, user: User=Depends(current_user), db: Session=Depends(get_db)):
    employee = db.query(Employee).filter(func.lower(Employee.email) == user.email.lower()).first()
    if not employee:
        raise HTTPException(404, "No employee profile is linked to this account")
    query = db.query(Entry, KPI).join(KPI, Entry.kpi_id == KPI.id).filter(Entry.employee_id == employee.id)
    if month:
        query = query.filter(Entry.month == month)
    return [{"id": entry.id, "employee_id": entry.employee_id, "kpi_id": entry.kpi_id,
             "month": entry.month, "actual": entry.actual, "performance_pct": entry.performance_pct,
             "score": entry.score, "final_score": entry.final_score, "status": entry.status,
             "employee_comments": entry.employee_comments, "supervisor_comments": entry.supervisor_comments,
             "kpi_name": kpi.name, "kpi_code": kpi.code, "target": entry.target_snapshot if entry.target_snapshot is not None else kpi.target,
             "unit_notes": kpi.unit_notes, "perspective": kpi.perspective,
             "has_context": entry.context is not None,
             "context_filename": _unpack_context(entry.context)[0] if entry.context else None}
            for entry, kpi in query.order_by(KPI.perspective, KPI.name).all()]


@router.post("/employees", response_model=EmployeeOut)
def add_employee(data: EmployeeCreate, db: Session=Depends(get_db), _=Depends(require_roles("Admin","Partner"))):
#def add_employee(data: EmployeeCreate, db: Session=Depends(get_db), _=Depends(current_user)):
    obj=Employee(**data.model_dump()); db.add(obj); db.commit(); db.refresh(obj); return obj

@router.delete("/admin/employees/{employee_id}", status_code=204)
def delete_employee(employee_id: int, db: Session=Depends(get_db), admin: User=Depends(require_roles("Admin"))):
    employee = db.get(Employee, employee_id)
    if not employee:
        raise HTTPException(404, "Employee not found")
    linked_user = None
    if employee.email:
        linked_user = db.query(User).filter(func.lower(User.email) == employee.email.lower()).first()
        if linked_user:
            if linked_user.id == admin.id:
                raise HTTPException(409, "Cannot delete the employee record linked to your admin account")
            if linked_user.role == "Admin" and linked_user.is_active:
                active_admins = db.query(User).filter(User.role == "Admin", User.is_active.is_(True)).count()
                if active_admins <= 1:
                    raise HTTPException(409, "At least one active Admin account must remain")
    employee.is_active = False
    if linked_user:
        db.delete(linked_user)
    db.commit()
    return Response(status_code=204)

@router.get("/entries", response_model=list[EntryOut])
def entries(month: str|None=None, employee_id: int|None=None, db: Session=Depends(get_db), user:User=Depends(current_user)):
    q=db.query(Entry, Employee.name, KPI).join(Employee, Entry.employee_id==Employee.id).join(KPI, Entry.kpi_id==KPI.id)
    if month: q=q.filter(Entry.month==month)
    if user.role == "Employee":
        employee = db.query(Employee).filter(func.lower(Employee.email) == user.email.lower()).first()
        if not employee: raise HTTPException(404, "No employee profile is linked to this account")
        q=q.filter(Entry.employee_id==employee.id)
    elif employee_id:
        q=q.filter(Entry.employee_id==employee_id)
    return [{"id": entry.id, "employee_id": entry.employee_id, "employee_name": employee_name,
             "kpi_id": entry.kpi_id, "kpi_name": kpi.name, "month": entry.month,
             "actual": entry.actual, "performance_pct": entry.performance_pct, "score": entry.score,
             "target": entry.target_snapshot if entry.target_snapshot is not None else kpi.target,
             "final_score": entry.final_score, "status": entry.status,
             "employee_comments": entry.employee_comments,
             "supervisor_comments": entry.supervisor_comments,
             "has_context": entry.context is not None,
             "context_filename": _unpack_context(entry.context)[0] if entry.context else None}
            for entry, employee_name, kpi in q.all()]

@router.post("/entries/{entry_id}/context")
def upload_entry_context(entry_id: int, file: UploadFile=File(...), db: Session=Depends(get_db), user: User=Depends(current_user)):
    entry = db.get(Entry, entry_id)
    if not entry:
        raise HTTPException(404, "Entry not found")
    _check_entry_access(entry, user, db)
    if entry.status not in ("Draft", "Rejected"):
        raise HTTPException(400, "Documents can only be added to draft or rejected entries")
    filename = (file.filename or "").replace("\\", "/").rsplit("/", 1)[-1].strip()
    extension = "." + filename.rsplit(".", 1)[-1].lower() if "." in filename else ""
    if not filename or extension not in ALLOWED_CONTEXT_EXTENSIONS:
        raise HTTPException(415, "Attach a PDF, Word, Excel, PowerPoint, TXT, CSV, PNG, or JPG document")
    content = file.file.read(MAX_CONTEXT_SIZE + 1)
    if not content:
        raise HTTPException(400, "The selected document is empty")
    if len(content) > MAX_CONTEXT_SIZE:
        raise HTTPException(413, "Documents must be 10 MB or smaller")
    content_type = mimetypes.guess_type(filename)[0] or "application/octet-stream"
    entry.context = _pack_context(filename, content_type, content)
    db.commit()
    return {"has_context": True, "context_filename": filename, "size": len(content)}

@router.get("/entries/{entry_id}/context")
def download_entry_context(entry_id: int, db: Session=Depends(get_db), user: User=Depends(current_user)):
    entry = db.get(Entry, entry_id)
    if not entry:
        raise HTTPException(404, "Entry not found")
    _check_entry_access(entry, user, db)
    if not entry.context:
        raise HTTPException(404, "No document is attached to this entry")
    filename, content_type, content = _unpack_context(entry.context)
    safe_filename = quote(filename, safe="")
    return Response(content=content, media_type=content_type,
                    headers={"Content-Disposition": f"attachment; filename*=UTF-8''{safe_filename}",
                             "X-Content-Type-Options": "nosniff"})

@router.delete("/entries/{entry_id}/context", status_code=204)
def delete_entry_context(entry_id: int, db: Session=Depends(get_db), user: User=Depends(current_user)):
    entry = db.get(Entry, entry_id)
    if not entry:
        raise HTTPException(404, "Entry not found")
    _check_entry_access(entry, user, db)
    if entry.status not in ("Draft", "Rejected"):
        raise HTTPException(400, "Documents can only be removed from draft or rejected entries")
    entry.context = None
    db.commit()

@router.patch("/entries/{entry_id}", response_model=EntryOut, response_model_exclude_unset=True)
def update_entry(entry_id:int, data:EntryUpdate, db:Session=Depends(get_db), user:User=Depends(current_user)):
    if user.role != "Employee":
        raise HTTPException(403, "Only employees can edit their KPI actuals")
    obj=db.get(Entry, entry_id)
    if not obj or obj.status not in ("Draft","Rejected"): raise HTTPException(400,"Entry unavailable for editing")
    employee = db.query(Employee).filter(func.lower(Employee.email) == user.email.lower()).first()
    if not employee or obj.employee_id != employee.id: raise HTTPException(404, "Entry not found")
    for k,v in data.model_dump(exclude_unset=True).items(): setattr(obj,k,v)
    calculate(obj, db.get(KPI,obj.kpi_id)); db.commit(); db.refresh(obj); return obj

@router.patch("/entries/{entry_id}/review-edit", response_model=EntryOut, response_model_exclude_unset=True)
def reviewer_edit_entry(entry_id: int, data: EntryUpdate, db: Session=Depends(get_db), _=Depends(require_roles("Admin"))):
    entry = db.get(Entry, entry_id)
    if not entry:
        raise HTTPException(404, "Entry not found")
    if entry.status != "Submitted":
        raise HTTPException(400, "Only entries awaiting review can be edited by a reviewer")
    changes = data.model_dump(exclude_unset=True)
    if not changes:
        raise HTTPException(400, "No entry changes provided")
    for field, value in changes.items():
        setattr(entry, field, value)
    calculate(entry, db.get(KPI, entry.kpi_id))
    db.commit()
    db.refresh(entry)
    return entry

@router.post("/entries/{entry_id}/submit", response_model=EntryOut, response_model_exclude_unset=True)
def submit(entry_id:int, db:Session=Depends(get_db), user:User=Depends(current_user)):
    if user.role != "Employee":
        raise HTTPException(403, "Only employees can submit their KPI actuals")
    obj=db.get(Entry,entry_id)
    if not obj or obj.actual is None: raise HTTPException(400,"Actual is required")
    employee = db.query(Employee).filter(func.lower(Employee.email) == user.email.lower()).first()
    if not employee or obj.employee_id != employee.id: raise HTTPException(404, "Entry not found")
    if obj.status not in ("Draft", "Rejected"): raise HTTPException(400, "Entry unavailable for submission")
    obj.status="Submitted"; db.commit(); db.refresh(obj); return obj

@router.patch("/entries/{entry_id}/review", response_model=EntryOut, response_model_exclude_unset=True)
def review(entry_id:int, data:ReviewUpdate, db:Session=Depends(get_db), _=Depends(require_roles("Admin"))):
    obj=db.get(Entry,entry_id)
    if not obj: raise HTTPException(404,"Entry not found")
    if obj.status != "Submitted": raise HTTPException(400, "Only entries submitted for review can be reviewed")
    if data.status == "Rejected" and not (data.supervisor_comments or "").strip():
        raise HTTPException(422, "Add reviewer feedback before rejecting an entry")
    if data.status == "Approved" and obj.actual is None:
        raise HTTPException(400, "An actual value is required before approval")
    obj.status=data.status; obj.supervisor_comments=data.supervisor_comments; db.commit(); db.refresh(obj); return obj

@router.post("/admin/generate-month")
def generate(data:GenerateRequest, db:Session=Depends(get_db), _=Depends(require_roles("Admin","Partner"))): return {"created":generate_month(db,data.month)}

@router.post("/entries/import-excel")
def import_entry_actuals(file: UploadFile=File(...), month: date=Form(...), employee_id: int|None=Form(default=None),
                         db: Session=Depends(get_db), user: User=Depends(current_user)):
    if user.role == "Employee":
        employee = db.query(Employee).filter(func.lower(Employee.email) == user.email.lower()).first()
        if not employee:
            raise HTTPException(404, "No employee profile is linked to this account")
    elif user.role == "Admin":
        if employee_id is None:
            raise HTTPException(422, "Select an employee for this import")
        employee = db.get(Employee, employee_id)
        if not employee:
            raise HTTPException(404, "Employee not found")
    else:
        raise HTTPException(403, "Only employees and admins can import KPI actuals")

    if not (file.filename or "").lower().endswith(".xlsx"):
        raise HTTPException(415, "Upload an Excel .xlsx workbook")
    content = file.file.read(MAX_IMPORT_SIZE + 1)
    if len(content) > MAX_IMPORT_SIZE:
        raise HTTPException(413, "Excel workbooks must be 10 MB or smaller")
    rows = _parse_excel_entry_rows(content)
    reporting_month = date(month.year, month.month, 1)
    assignments = (db.query(EmployeeKPI, KPI).join(KPI, EmployeeKPI.kpi_id == KPI.id)
                   .filter(EmployeeKPI.employee_id == employee.id, EmployeeKPI.is_active.is_(True)).all())
    if not assignments:
        raise HTTPException(400, "This employee has no active KPI assignments")

    issues = []
    resolved = []
    seen_kpis = set()
    for row in rows:
        excel_row = row.get("row_number")
        kpi = _match_import_kpi(row, assignments)
        if not kpi:
            issues.append({"row": excel_row, "kpi": row.get("label") or row.get("code"), "reason": "Could not match to an active employee KPI"})
            continue
        if kpi.id in seen_kpis:
            issues.append({"row": excel_row, "kpi": kpi.name, "reason": "Duplicate KPI row; the first matching row was kept"})
            continue
        seen_kpis.add(kpi.id)
        try:
            actual = float(row["actual"])
        except (TypeError, ValueError):
            issues.append({"row": excel_row, "kpi": kpi.name, "reason": "Actual must be a number"})
            continue
        if not math.isfinite(actual):
            issues.append({"row": excel_row, "kpi": kpi.name, "reason": "Actual must be a finite number"})
            continue
        resolved.append((row, kpi, actual))

    created_entries = 0
    for assignment, kpi in assignments:
        entry = db.query(Entry).filter_by(employee_id=employee.id, kpi_id=kpi.id, month=reporting_month).first()
        if entry is None:
            entry = Entry(employee_id=employee.id, kpi_id=kpi.id, month=reporting_month,
                          target_snapshot=assignment.target, status="Draft")
            db.add(entry)
            db.flush()
            created_entries += 1

    imported = 0
    protected = 0
    for row, kpi, actual in resolved:
        entry = db.query(Entry).filter_by(employee_id=employee.id, kpi_id=kpi.id, month=reporting_month).first()
        if entry.status not in ("Draft", "Rejected"):
            protected += 1
            issues.append({"row": None, "kpi": kpi.name, "reason": f"Entry is {entry.status}; existing reviewed data was left unchanged"})
            continue
        if entry.target_snapshot is None:
            entry.target_snapshot = next(assignment.target for assignment, assigned_kpi in assignments if assigned_kpi.id == kpi.id)
        entry.actual = actual
        entry.employee_comments = row["comments"]
        calculate(entry, kpi)
        imported += 1

    db.commit()
    return {"employee_id": employee.id, "employee_name": employee.name, "month": reporting_month.isoformat(),
            "imported": imported, "created_entries": created_entries, "protected": protected, "issues": issues}

@router.get("/dashboard")
def dashboard(month:str, db:Session=Depends(get_db), _=Depends(require_roles("Manager","Partner","Admin"))):
    team_rows=(db.query(Employee.team,func.avg(Entry.final_score)).join(Entry).filter(Entry.month==month).group_by(Employee.team).all())
    people=(db.query(Employee.id, Employee.name, Employee.team, func.sum(Entry.final_score))
            .join(Entry).filter(Entry.month==month).group_by(Employee.id).all())
    return {"month":month,"teams":[{"team":x,"average_score":y or 0} for x,y in team_rows],"employees":[{"employee_id":employee_id,"employee":name,"team":team,"total_score":total_score or 0,"percentage_score":(total_score or 0) / 5 * 100} for employee_id,name,team,total_score in people]}

@router.get("/dashboard/employee-trend")
def employee_trend(employee_id:int, year:int, db:Session=Depends(get_db), _=Depends(require_roles("Manager","Partner","Admin"))):
    employee = db.get(Employee, employee_id)
    if not employee: raise HTTPException(404, "Employee not found")
    start, end = date(year, 1, 1), date(year + 1, 1, 1)
    rows = (db.query(Entry.month, func.sum(Entry.final_score))
              .filter(Entry.employee_id == employee_id, Entry.month >= start, Entry.month < end)
              .group_by(Entry.month).order_by(Entry.month).all())
    percentage_by_month = {month.month: (total_score or 0) / 5 * 100 for month, total_score in rows}
    return {"employee_id":employee.id,"employee":employee.name,"year":year,
            "points":[{"month":date(year, month, 1).isoformat(), "percentage_score":percentage_by_month.get(month, 0)} for month in range(1, 13)]}
